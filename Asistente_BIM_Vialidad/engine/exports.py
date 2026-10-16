# -*- coding: utf-8 -*-
"""Exportaciones reales: XLSX (con fórmulas) y BCF 2.1."""
import io
import uuid
import math
import zipfile
import datetime
from xml.sax.saxutils import escape

DISC = ["GEO", "PAV", "DRE", "EST", "SEG", "SER"]


def to_xlsx(result, project_name, rules):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = PatternFill("solid", fgColor="1E293B")
    hfont = Font(bold=True, color="FFFFFF")

    def header(ws, row, cols):
        for c, t in enumerate(cols, 1):
            cell = ws.cell(row=row, column=c, value=t)
            cell.fill, cell.font = head, hfont
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    # --- Interferencias (hoja base de las fórmulas)
    wi = wb.active
    wi.title = "Interferencias"
    cols = ["ID", "Pk (m)", "Pk", "Tramo", "Tipo", "Esp. A", "Esp. B", "Archivo A", "GUID A", "STEP A",
            "Archivo B", "GUID B", "STEP B", "Distancia (m)", "Severidad", "Responsable", "Estado",
            "Fecha compromiso", "Descripción y criterio", "X", "Y", "Z"]
    header(wi, 1, cols)
    for r, it in enumerate(result["issues"], 2):
        row = [it["id"], None if it["pk"] is None else round(it["pk"], 2), it["pk_text"], it["level"], it["kind"],
               it["specs"][0], it["specs"][1], it["fileA"], it["guidA"], it["stepA"], it["fileB"], it["guidB"],
               it["stepB"], it["distance"], it["severity"], it["owner"], it["status"], "", it["desc"],
               it["point"][0], it["point"][1], it["point"][2]]
        for c, v in enumerate(row, 1):
            wi.cell(row=r, column=c, value=v)
    n = max(2, len(result["issues"]) + 1)
    for i, w in enumerate([11, 10, 16, 22, 9, 7, 7, 34, 24, 9, 34, 24, 9, 11, 11, 30, 12, 14, 90, 12, 12, 8], 1):
        wi.column_dimensions[get_column_letter(i)].width = w
    wi.freeze_panes = "A2"
    wi.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"

    # --- Resumen
    ws = wb.create_sheet("Resumen", 0)
    ws["A1"] = f"Libro de Coordinación Vial — {project_name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generado: {datetime.datetime.now():%Y-%m-%d %H:%M}"
    rng = f"Interferencias!$A$2:$A${n}"
    sev = f"Interferencias!$O$2:$O${n}"
    est = f"Interferencias!$Q$2:$Q${n}"
    typ = f"Interferencias!$E$2:$E${n}"
    rows = [
        ("Interferencias totales", f"=COUNTA({rng})"),
        ("Activas (no resueltas)", f'=COUNTA({rng})-COUNTIF({est},"Resuelta")'),
        ("Resueltas", f'=COUNTIF({est},"Resuelta")'),
        ("% resuelto", f'=IF(B4=0,0,B6/B4)'),
        ("Críticas activas", f'=COUNTIFS({sev},"Crítica",{est},"<>Resuelta")'),
        ("Choques duros", f'=COUNTIF({typ},"choque")'),
        ("Infracciones de holgura", f'=COUNTIF({typ},"holgura")'),
        ("Elementos auditados", result["totalElements"]),
        ("Completitud media (%)", result["kpi"]["avgCompleteness"]),
        ("Proxy medio (%)", result["kpi"]["avgProxy"]),
        ("Eje para Pk", result["axis"]["source"]),
        ("Pares de elementos evaluados", result["pairs_tested"]),
    ]
    for i, (k, v) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws["B7"].number_format = "0%"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60

    # --- Matriz (COUNTIFS vivos)
    wm = wb.create_sheet("Matriz", 1)
    header(wm, 1, ["Esp. \\ Esp."] + DISC)
    ca, cb = "Interferencias!$F$2:$F$%d" % n, "Interferencias!$G$2:$G$%d" % n
    for r, d in enumerate(DISC, 2):
        wm.cell(row=r, column=1, value=d).font = Font(bold=True)
        for c, e in enumerate(DISC, 2):
            col = get_column_letter(c)
            if d == e:
                wm.cell(row=r, column=c, value=f'=COUNTIFS({ca},$A{r},{cb},{col}$1)')
            else:
                wm.cell(row=r, column=c, value=f'=COUNTIFS({ca},$A{r},{cb},{col}$1)+COUNTIFS({ca},{col}$1,{cb},$A{r})')

    # --- Modelos
    wo = wb.create_sheet("Modelos")
    header(wo, 1, ["Especialidad", "Archivo", "Esquema", "Elementos", "Completitud %", "Psets %", "Geometría %",
                   "Nombre %", "Clasificación %", "Proxy %", "NDI (est.)", "CRS", "Origen especialidad"])
    for r, m in enumerate(result["models"], 2):
        g = m.get("georef") or {}
        for c, v in enumerate([m["code"], m["file"], m["format"], m["elements"], m["completitud"], m["pset"],
                               m["geom_pct"], m["name_pct"], m["class_pct"], m["proxy"], m["lod"],
                               g.get("crs") or "no definido", m["disc_src"]], 1):
            wo.cell(row=r, column=c, value=v)
    for i, w in enumerate([12, 60, 9, 11, 13, 9, 12, 9, 14, 9, 13, 16, 22], 1):
        wo.column_dimensions[get_column_letter(i)].width = w

    # --- Reglas
    wr = wb.create_sheet("Reglas")
    header(wr, 1, ["Par", "Holgura (m)", "Severidad choque", "Norma", "Verificado", "Descripción"])
    d = dict(rules.get("pares", {}))
    for r, (k, v) in enumerate(sorted(d.items()), 2):
        for c, x in enumerate([k, v.get("holgura_m"), v.get("choque"), v.get("norma"), "sí" if v.get("verificado") else "no",
                               v.get("descripcion")], 1):
            wr.cell(row=r, column=c, value=x)
    wr.cell(row=len(d) + 3, column=1, value=rules.get("nota", ""))

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _camera(point, dist=18.0):
    x, y, z = point
    off = dist / math.sqrt(3)
    cam = (x + off, y - off, z + off)
    d = (-off, off, -off)
    n = math.sqrt(sum(c * c for c in d))
    d = tuple(c / n for c in d)
    return cam, d


def to_bcf(result, project_name, author="coordinacion@bim"):
    """BCF 2.1 con un tópico + viewpoint por interferencia."""
    bio = io.BytesIO()
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    prio = {"Crítica": "Alta", "Mayor": "Media", "Moderada": "Baja"}
    status = {"Activa": "Active", "En Proceso": "In Progress", "Resuelta": "Resolved"}
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("bcf.version", '<?xml version="1.0" encoding="UTF-8"?>\n<Version VersionId="2.1" '
                   'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:xsd="http://www.w3.org/2001/XMLSchema">'
                   '<DetailedVersion>2.1</DetailedVersion></Version>')
        for it in result["issues"]:
            g = str(uuid.uuid5(uuid.NAMESPACE_URL, it["key"]))
            vp = str(uuid.uuid4())
            title = f"{it['id']} {it['pair']} {it['pk_text']} ({it['kind']})"
            markup = f'''<?xml version="1.0" encoding="UTF-8"?>
<Markup xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:xsd="http://www.w3.org/2001/XMLSchema">
  <Header>
    <File IfcProject="" isExternal="false"><Filename>{escape(it['fileA'])}</Filename></File>
    <File IfcProject="" isExternal="false"><Filename>{escape(it['fileB'])}</Filename></File>
  </Header>
  <Topic Guid="{g}" TopicType="Clash" TopicStatus="{status.get(it['status'], 'Active')}">
    <Title>{escape(title)}</Title>
    <Priority>{prio.get(it['severity'], 'Media')}</Priority>
    <Labels>{escape(project_name)}</Labels>
    <CreationDate>{now}</CreationDate>
    <CreationAuthor>{escape(author)}</CreationAuthor>
    <AssignedTo>{escape(it['owner'])}</AssignedTo>
    <Description>{escape(it['desc'])}</Description>
  </Topic>
  <Viewpoints Guid="{vp}"><Viewpoint>viewpoint.bcfv</Viewpoint></Viewpoints>
</Markup>'''
            cam, d = _camera(it["point"])
            view = f'''<?xml version="1.0" encoding="UTF-8"?>
<VisualizationInfo Guid="{vp}">
  <Components>
    <Selection>
      <Component IfcGuid="{it['guidA']}"/>
      <Component IfcGuid="{it['guidB']}"/>
    </Selection>
    <Visibility DefaultVisibility="true"/>
  </Components>
  <PerspectiveCamera>
    <CameraViewPoint><X>{cam[0]:.3f}</X><Y>{cam[1]:.3f}</Y><Z>{cam[2]:.3f}</Z></CameraViewPoint>
    <CameraDirection><X>{d[0]:.5f}</X><Y>{d[1]:.5f}</Y><Z>{d[2]:.5f}</Z></CameraDirection>
    <CameraUpVector><X>0</X><Y>0</Y><Z>1</Z></CameraUpVector>
    <FieldOfView>60</FieldOfView>
  </PerspectiveCamera>
</VisualizationInfo>'''
            z.writestr(f"{g}/markup.bcf", markup)
            z.writestr(f"{g}/viewpoint.bcfv", view)
    return bio.getvalue()
