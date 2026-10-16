import sys
import subprocess
import os

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    install('python-docx')
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_letter():
    doc = Document()
    
    # Add logo
    logo_path = r"C:/Users/erick.rodriguez/.gemini/antigravity/brain/18f6f20f-9ae4-494e-83df-686f73b41bab/.user_uploaded/media_1791458756918_8e27cfb2.png"
    try:
        if os.path.exists(logo_path):
            doc.add_picture(logo_path, width=Inches(2.5))
            last_paragraph = doc.paragraphs[-1] 
            last_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    except Exception as e:
        print(f"Could not add logo: {e}")

    # Add Date
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = date_p.add_run("Fecha: 08 de Octubre de 2026")
    
    doc.add_paragraph() # Spacing
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("CARTA DE RECOMENDACIÓN")
    title_run.bold = True
    title_run.font.size = Pt(14)
    
    doc.add_paragraph() # Spacing
    
    # Body
    body_text = (
        "A quien corresponda,\n\n"
        "Por medio de la presente, tengo el agrado de recomendar ampliamente al Sr. Carlos Daniel Toro Mattei, "
        "portador del RUT 25950828-2.\n\n"
        "El Sr. Toro prestó sus servicios en nuestra empresa, AERIAL SECURITY, durante el período comprendido "
        "entre los años 2018 y 2019, desempeñándose de manera sobresaliente en el cargo de Instalador de Estructuras "
        "Metálicas en Altura.\n\n"
        "Durante su permanencia con nosotros, demostró ser un trabajador sumamente proactivo, altamente colaborativo "
        "con sus compañeros de equipo y con una gran eficiencia en el desarrollo de sus labores diarias. Además, "
        "cumplió a cabalidad con los más altos y estrictos estándares de seguridad exigidos para los trabajos de riesgo y en altura.\n\n"
        "Su disposición para resolver problemas, su capacidad para trabajar en equipo y su compromiso ético y profesional "
        "lo convierten en un elemento de gran valor para cualquier organización.\n\n"
        "Es por ello que no dudo en extender esta recomendación, con la certeza de que el Sr. Carlos Daniel Toro Mattei "
        "sabrá desempeñarse con el mismo nivel de responsabilidad, dedicación y profesionalismo que demostró durante "
        "su tiempo de servicio con nosotros.\n\n"
        "Quedo a su entera disposición para cualquier información adicional que requieran sobre su desempeño."
    )
    
    p = doc.add_paragraph(body_text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph() # Spacing
    doc.add_paragraph() # Spacing
    
    # Sign-off
    sign_p = doc.add_paragraph()
    sign_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sign_p.add_run("Atentamente,\n\n\n\n")
    sign_p.add_run("___________________________________\n")
    sign_p.add_run("Firma / Sello\n")
    
    bold_company = sign_p.add_run("AERIAL SECURITY\n")
    bold_company.bold = True
    
    sign_p.add_run("Fabricación de Estructuras Industriales")
    
    output_path = r"c:\OD\OneDrive - Ingeniería Cuatro S.A\Documentos\_00.PRUEBAS.DYNAMO\ICuatro_Dev_BIM\Carta_Recomendacion_Carlos_Toro.docx"
    doc.save(output_path)
    print(f"Document saved to {output_path}")

if __name__ == '__main__':
    create_letter()

